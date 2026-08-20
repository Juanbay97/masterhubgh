"""Servicio de cruce de selección: contrato de columnas + resolutores de fila.

Fuente única de verdad para el orden/etiquetas de las 49 columnas del cruce
(bandeja + export Excel, PR3) y para la resolución de cada campo por
candidato. También expone el gate de lectura HR-READ reutilizado por el
snapshot de afiliación/contratación (`contratacion_service.affiliation_contract_snapshot`).

PR3 agrega los endpoints whitelisted de bandeja (`list_cruce_candidates`) y
export Excel (`export_cruce_xlsx`), ambos gateados por `validate_hr_or_selection_read_access`
ANTES de cualquier lectura a base de datos.
"""

import base64
import unicodedata
from collections import namedtuple
from datetime import date
from io import BytesIO

import frappe
from frappe.utils import getdate

from hubgh.hubgh.candidate_states import is_candidate_status
from hubgh.hubgh.display_labels import (
	get_punto_name_map,
	resolve_candidate_location_labels,
	resolve_catalog_display_name,
)
from hubgh.hubgh.role_matrix import user_has_any_role


ColumnSpec = namedtuple("ColumnSpec", ["key", "label"])

# Fuente única de verdad: 49 columnas ordenadas del cruce. `key` es único
# (incluso cuando `label` se repite entre grupos one-hot, p.ej. XS/S/M/L/XL
# para camisa/pantalón/delantal).
CRUCE_COLUMNS = (
	ColumnSpec("tipo_documento", "T.D"),
	ColumnSpec("nombre", "NOMBRE"),
	ColumnSpec("edad", "EDAD"),
	ColumnSpec("numero_documento", "CC"),
	ColumnSpec("fecha_lugar_nacimiento", "FECHA Y LUGAR DE NACIMIENTO"),
	ColumnSpec("direccion", "DIRECCION"),
	ColumnSpec("ciudad_barrio", "CIUDAD Y BARRIO"),
	ColumnSpec("telefono_fijo", "TELEFONO FIJO"),
	ColumnSpec("numero_celular", "NUMERO CELULAR"),
	ColumnSpec("correo_electronico", "CORREO ELECTRONICO"),
	ColumnSpec("rh", "RH"),
	ColumnSpec("estado_civil__soltero", "SOLTERO"),
	ColumnSpec("estado_civil__casado", "CASADO"),
	ColumnSpec("estado_civil__ulibre", "ULIBRE"),
	ColumnSpec("estado_civil__viudo", "VIUDO"),
	ColumnSpec("estado_civil__divorciado", "DIVORCIADO"),
	ColumnSpec("eps", "EPS"),
	ColumnSpec("nivel_educativo__primaria", "PRIMARIA"),
	ColumnSpec("nivel_educativo__bachiller_clasico", "BACHILLER CLASICO"),
	ColumnSpec("nivel_educativo__bachiller_tecnico", "BACHILLER TECNICO"),
	ColumnSpec("nivel_educativo__tecnico", "TECNICO"),
	ColumnSpec("nivel_educativo__tecnologo", "TECNOLOGO"),
	ColumnSpec("nivel_educativo__universitario", "UNIVERSITARIO"),
	ColumnSpec("nivel_educativo__postgrado", "POSTGRADO"),
	ColumnSpec("afp", "AFP"),
	ColumnSpec("emergencia", "EMERGENCIA"),
	ColumnSpec("celular_emergencia", "CELULAR"),
	ColumnSpec("alergico", "ALERGICO"),
	ColumnSpec("talla_camisa__xs", "XS"),
	ColumnSpec("talla_camisa__s", "S"),
	ColumnSpec("talla_camisa__m", "M"),
	ColumnSpec("talla_camisa__l", "L"),
	ColumnSpec("talla_camisa__xl", "XL"),
	ColumnSpec("talla_pantalon__xs", "XS"),
	ColumnSpec("talla_pantalon__s", "S"),
	ColumnSpec("talla_pantalon__m", "M"),
	ColumnSpec("talla_pantalon__l", "L"),
	ColumnSpec("talla_pantalon__xl", "XL"),
	ColumnSpec("zapatos", "ZAPATOS"),
	ColumnSpec("talla_delantal__xs", "XS"),
	ColumnSpec("talla_delantal__s", "S"),
	ColumnSpec("talla_delantal__m", "M"),
	ColumnSpec("talla_delantal__l", "L"),
	ColumnSpec("talla_delantal__xl", "XL"),
	ColumnSpec("ciudad_expedicion", "CIUDAD DE EXPEDICION"),
	ColumnSpec("cesantias", "CESANTIAS"),
	ColumnSpec("banco", "BANCO"),
	ColumnSpec("numero_cuenta", "NUMERO DE CUENTA"),
	ColumnSpec("ciudad_residencia", "CIUDAD"),
)

# Los 6 roles HR-READ autorizados a leer el cruce (Selección + HR-EXT + RRLL, solo lectura).
CRUCE_ROLES = (
	"HR Selection",
	"Gestión Humana",
	"GH - Bandeja General",
	"Gerente GH",
	"HR Labor Relations",
	"Relaciones Laborales Jefe",
)

# Mapa congelado estado_civil → columna one-hot. Comparación por texto normalizado
# (sin tildes, minúsculas) contra las opciones vivas de Candidato.estado_civil.
ESTADO_CIVIL_ONEHOT = (
	("soltero", "estado_civil__soltero"),
	("casado", "estado_civil__casado"),
	("union libre", "estado_civil__ulibre"),
	("viudo", "estado_civil__viudo"),
	("divorciado", "estado_civil__divorciado"),
)

# Mapa congelado (7-tupla) nivel educativo → columna one-hot, por coincidencia de
# palabra clave normalizada contra la `description` del catálogo Nivel Educativo Siesa.
#
# REGLA APROBADA POR NEGOCIO (decisión vinculante del usuario, ya no un supuesto):
# "todo bachiller" → BACHILLER CLASICO por defecto; BACHILLER TECNICO queda vacío
# a menos que el catálogo diga EXPLÍCITAMENTE "bachiller técnico". "TÉCNICO LABORAL"
# NO es una designación explícita de bachiller-técnico, por lo tanto NO mapea a
# BACHILLER TECNICO (ver `_nivel_educativo_onehot`: el match de BACHILLER TECNICO
# corre con prioridad/antes que el catch-all genérico de "bachiller" para que una
# descripción explícita de bachiller técnico no caiga en CLASICO). PREESCOLAR/SIN
# DEFINIR/OTROS quedan intencionalmente sin columna (grupo en blanco).
NIVEL_EDUCATIVO_ONEHOT = (
	("PRIMARIA", "nivel_educativo__primaria", ("primaria",)),
	("BACHILLER CLASICO", "nivel_educativo__bachiller_clasico", ("secundaria", "media", "bachiller clasico", "bachiller")),
	("BACHILLER TECNICO", "nivel_educativo__bachiller_tecnico", ("bachiller tecnico",)),
	("TECNICO", "nivel_educativo__tecnico", ("tecnica profesional", "tecnico profesional")),
	("TECNOLOGO", "nivel_educativo__tecnologo", ("tecnologica", "tecnologo")),
	("UNIVERSITARIO", "nivel_educativo__universitario", ("universitaria", "universitario")),
	("POSTGRADO", "nivel_educativo__postgrado", ("especializacion", "maestria", "doctorado", "postgrado")),
)

_SIZE_KEYS = ("xs", "s", "m", "l", "xl")


def _first_value(*values):
	for value in values:
		if value not in (None, ""):
			return value
	return None


def _get(doc, field):
	return doc.get(field) if doc is not None else None


def _normalize_match_text(value):
	if value is None:
		return ""
	raw = str(value).strip().lower().replace("-", " ")
	raw = "".join(ch for ch in unicodedata.normalize("NFKD", raw) if not unicodedata.combining(ch))
	return " ".join(raw.split())


def resolve_shared_field(datos, candidato, fieldname):
	"""Resuelve un campo presente en ambos DocTypes (Datos Contratacion gana sobre Candidato).

	Única fuente de verdad reutilizada por `build_cruce_row` y por el bloque
	demográficos/dotación del snapshot de afiliación, para evitar drift.
	"""
	return _first_value(_get(datos, fieldname), _get(candidato, fieldname))


def compute_edad(fecha_nacimiento, as_of=None):
	"""Edad por aritmética de calendario real (no `date_diff // 365`)."""
	if not fecha_nacimiento:
		return None
	birth = getdate(fecha_nacimiento)
	reference = getdate(as_of) if as_of else date.today()
	edad = reference.year - birth.year
	if (reference.month, reference.day) < (birth.month, birth.day):
		edad -= 1
	return edad


def _has_cruce_read_access(user=None):
	user = user or frappe.session.user
	return user == "Administrator" or user_has_any_role(user, *CRUCE_ROLES)


def validate_hr_or_selection_read_access():
	"""Gate de solo lectura para HR-READ (HR-EXT + RRLL). Corre ANTES de cualquier lectura DB."""
	if not _has_cruce_read_access():
		frappe.throw("No autorizado para consultar el cruce de selección.")


def _size_onehot(prefix, value):
	keys = [f"{prefix}__{k}" for k in _SIZE_KEYS]
	blank = dict.fromkeys(keys, "")
	normalized = str(value or "").strip().upper()
	for size in _SIZE_KEYS:
		if normalized == size.upper():
			blank[f"{prefix}__{size}"] = "X"
			break
	return blank


def _estado_civil_onehot(value):
	keys = [key for _, key in ESTADO_CIVIL_ONEHOT]
	blank = dict.fromkeys(keys, "")
	normalized = _normalize_match_text(value)
	if not normalized:
		return blank
	for match_text, key in ESTADO_CIVIL_ONEHOT:
		if normalized == match_text:
			blank[key] = "X"
			break
	return blank


def _nivel_educativo_onehot(nivel_educativo_siesa):
	keys = [key for _, key, _ in NIVEL_EDUCATIVO_ONEHOT]
	blank = dict.fromkeys(keys, "")
	if not nivel_educativo_siesa:
		return blank
	description = resolve_catalog_display_name("Nivel Educativo Siesa", nivel_educativo_siesa)
	normalized = _normalize_match_text(description)
	if not normalized:
		return blank

	# BACHILLER TECNICO se evalúa PRIMERO y por separado: es un caso explícito y más
	# específico que el catch-all genérico "bachiller" de BACHILLER CLASICO. Si se
	# revisara en el orden de la tupla (CLASICO antes que TECNICO), una descripción
	# como "BACHILLER TÉCNICO" caería incorrectamente en CLASICO por el match genérico.
	bachiller_tecnico_keywords = next(
		kw for label, _key, kw in NIVEL_EDUCATIVO_ONEHOT if label == "BACHILLER TECNICO"
	)
	if any(keyword in normalized for keyword in bachiller_tecnico_keywords):
		blank["nivel_educativo__bachiller_tecnico"] = "X"
		return blank

	for _label, key, keywords in NIVEL_EDUCATIVO_ONEHOT:
		if key == "nivel_educativo__bachiller_tecnico":
			continue
		if any(keyword in normalized for keyword in keywords):
			blank[key] = "X"
			break
	return blank


def _build_alergico(candidato):
	if not int(_get(candidato, "tiene_alergias") or 0):
		return "NO"
	descripcion = str(_get(candidato, "descripcion_alergias") or "").strip()
	return descripcion or "SI"


def _build_fecha_lugar_nacimiento(fecha_nacimiento, datos):
	"""Fecha + lugar de nacimiento; el lugar SOLO viene de Datos Contratacion
	(`*_nacimiento_siesa`), sin fallback a `procedencia_*`. Blanco cuando falte."""
	fecha_str = str(fecha_nacimiento) if fecha_nacimiento else ""
	lugar = ""
	ciudad_code = _get(datos, "ciudad_nacimiento_siesa")
	departamento_code = _get(datos, "departamento_nacimiento_siesa")
	pais_code = _get(datos, "pais_nacimiento_siesa")
	if ciudad_code or departamento_code or pais_code:
		labels = resolve_candidate_location_labels(pais=pais_code, departamento=departamento_code, ciudad=ciudad_code)
		lugar = ", ".join(p for p in [labels.get("ciudad"), labels.get("departamento"), labels.get("pais")] if p)
	if fecha_str and lugar:
		return f"{fecha_str} - {lugar}"
	return fecha_str or lugar


def _resolve_full_name(candidato, datos):
	nombres = resolve_shared_field(datos, candidato, "nombres") or ""
	apellidos = resolve_shared_field(datos, candidato, "apellidos")
	if not apellidos:
		primer = resolve_shared_field(datos, candidato, "primer_apellido") or ""
		segundo = resolve_shared_field(datos, candidato, "segundo_apellido") or ""
		apellidos = " ".join(p for p in [primer, segundo] if p)
	return " ".join(p for p in [nombres, apellidos] if p).strip()


def build_cruce_row(candidato, datos=None):
	"""Construye una fila del cruce (dict keyed por `ColumnSpec.key`).

	`candidato` es un Document/`_dict` de Candidato (obligatorio). `datos` es un
	Document/`_dict` de Datos Contratacion o `None`. EDAD se calcula al vuelo,
	nunca se persiste.
	"""
	fecha_nacimiento = resolve_shared_field(datos, candidato, "fecha_nacimiento")
	ciudad_residencia_code = _first_value(
		_get(datos, "ciudad_residencia_siesa"), _get(datos, "ciudad"), _get(candidato, "ciudad")
	)
	ciudad_labels = resolve_candidate_location_labels(ciudad=ciudad_residencia_code)
	ciudad_residencia = ciudad_labels.get("ciudad") or ciudad_residencia_code or ""
	barrio = resolve_shared_field(datos, candidato, "barrio") or ""

	row = {
		"tipo_documento": resolve_shared_field(datos, candidato, "tipo_documento") or "",
		"nombre": _resolve_full_name(candidato, datos),
		"edad": compute_edad(fecha_nacimiento) if fecha_nacimiento else "",
		"numero_documento": resolve_shared_field(datos, candidato, "numero_documento") or "",
		"fecha_lugar_nacimiento": _build_fecha_lugar_nacimiento(fecha_nacimiento, datos),
		"direccion": resolve_shared_field(datos, candidato, "direccion") or "",
		"ciudad_barrio": " - ".join(p for p in [ciudad_residencia, barrio] if p),
		"telefono_fijo": _get(candidato, "telefono_fijo") or "",
		"numero_celular": resolve_shared_field(datos, candidato, "celular") or "",
		"correo_electronico": resolve_shared_field(datos, candidato, "email") or "",
		"rh": _get(candidato, "grupo_sanguineo") or "",
		"eps": resolve_shared_field(datos, candidato, "eps_siesa") or "",
		"afp": resolve_shared_field(datos, candidato, "afp_siesa") or "",
		"emergencia": _get(candidato, "contacto_emergencia_nombre") or "",
		"celular_emergencia": _get(candidato, "contacto_emergencia_telefono") or "",
		"alergico": _build_alergico(candidato),
		"zapatos": _get(candidato, "numero_zapatos") or "",
		# ciudad_expedicion_siesa vive SOLO en Datos Contratacion — sin fallback a Candidato.
		"ciudad_expedicion": _get(datos, "ciudad_expedicion_siesa") or "",
		"cesantias": resolve_shared_field(datos, candidato, "cesantias_siesa") or "",
		"banco": resolve_shared_field(datos, candidato, "banco_siesa") or "",
		"numero_cuenta": resolve_shared_field(datos, candidato, "numero_cuenta_bancaria") or "",
		"ciudad_residencia": ciudad_residencia,
	}
	row.update(_estado_civil_onehot(resolve_shared_field(datos, candidato, "estado_civil")))
	row.update(_nivel_educativo_onehot(resolve_shared_field(datos, candidato, "nivel_educativo_siesa")))
	row.update(_size_onehot("talla_camisa", _get(candidato, "talla_camisa")))
	row.update(_size_onehot("talla_pantalon", _get(candidato, "talla_pantalon")))
	row.update(_size_onehot("talla_delantal", _get(candidato, "talla_delantal")))
	return row


# ---------------------------------------------------------------------------
# Bandeja del cruce (PR3) — listado + export Excel
# ---------------------------------------------------------------------------


def _base_cruce_filters(pdv=None, fecha_desde=None, fecha_hasta=None, search=None):
	filters = {}
	if pdv:
		filters["pdv_destino"] = pdv
	if fecha_desde and fecha_hasta:
		filters["fecha_tentativa_ingreso"] = ["between", [fecha_desde, fecha_hasta]]
	elif fecha_desde:
		filters["fecha_tentativa_ingreso"] = [">=", fecha_desde]
	elif fecha_hasta:
		filters["fecha_tentativa_ingreso"] = ["<=", fecha_hasta]
	if search:
		filters["numero_documento"] = ["like", f"%{search}%"]
	return filters


def _query_cruce_candidatos(estado=None, pdv=None, fecha_desde=None, fecha_hasta=None, search=None):
	"""Query compartida por `list_cruce_candidates` y `export_cruce_xlsx`.

	SIEMPRE excluye Rechazado, sin importar los filtros recibidos (incluso si
	`estado` pide explícitamente "Rechazado", el resultado queda vacío para ese caso).
	"""
	filters = _base_cruce_filters(pdv=pdv, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta, search=search)
	rows = frappe.get_all("Candidato", filters=filters, fields=["*"], order_by="creation desc")
	rows = [row for row in rows if not is_candidate_status(row.get("estado_proceso"), "Rechazado")]
	if estado:
		rows = [row for row in rows if is_candidate_status(row.get("estado_proceso"), estado)]
	return rows


def _attach_datos_contratacion(candidatos):
	"""Mapa `candidato -> Datos Contratacion` (dict) para los candidatos dados."""
	names = [row.get("name") for row in candidatos if row.get("name")]
	if not names:
		return {}
	rows = frappe.get_all("Datos Contratacion", filters={"candidato": ["in", names]}, fields=["*"])
	return {row.get("candidato"): row for row in rows}


@frappe.whitelist()
def list_cruce_candidates(estado=None, pdv=None, fecha_desde=None, fecha_hasta=None, search=None):
	"""Bandeja de lectura del cruce de selección (HR-READ, solo lectura).

	Filtra por estado_proceso/pdv_destino/rango de fecha_tentativa_ingreso/búsqueda
	por número de documento. SIEMPRE excluye Rechazado. Resultado vacío -> lista vacía,
	sin error.
	"""
	validate_hr_or_selection_read_access()
	rows = _query_cruce_candidatos(
		estado=estado, pdv=pdv, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta, search=search
	)
	pdv_name_map = get_punto_name_map([row.get("pdv_destino") for row in rows])
	return [
		{
			"name": row.get("name"),
			"full_name": _resolve_full_name(row, None),
			"numero_documento": row.get("numero_documento") or "",
			"pdv_destino": row.get("pdv_destino"),
			"pdv_destino_nombre": pdv_name_map.get(row.get("pdv_destino"), row.get("pdv_destino") or ""),
			"cargo_postulado": row.get("cargo_postulado"),
			"estado_proceso": row.get("estado_proceso"),
			"fecha_tentativa_ingreso": row.get("fecha_tentativa_ingreso"),
			"creation": row.get("creation"),
		}
		for row in rows
	]


@frappe.whitelist()
def export_cruce_xlsx(filters=None):
	"""Exporta el cruce de selección a Excel: 49 columnas fijas (`CRUCE_COLUMNS`).

	Gate HR-READ ANTES de cualquier lectura DB. Auditado (usuario, timestamp,
	cantidad de candidatos) en cada llamada, incluso cuando `count == 0`
	(se retorna un workbook solo-encabezados y de todas formas se audita).

	Returns:
		{"filename": "...", "content_b64": "<base64>", "count": N}
	"""
	validate_hr_or_selection_read_access()

	import openpyxl
	from openpyxl.styles import Alignment, Font, PatternFill

	filters = frappe.parse_json(filters) if filters else {}
	rows = _query_cruce_candidatos(
		estado=filters.get("estado"),
		pdv=filters.get("pdv"),
		fecha_desde=filters.get("fecha_desde"),
		fecha_hasta=filters.get("fecha_hasta"),
		search=filters.get("search"),
	)
	datos_by_candidato = _attach_datos_contratacion(rows)

	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = "Cruce Selección"

	headers = [column.label for column in CRUCE_COLUMNS]
	ws.append(headers)

	header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
	header_font = Font(bold=True, color="FFFFFF")
	for col_idx in range(1, len(headers) + 1):
		cell = ws.cell(row=1, column=col_idx)
		cell.fill = header_fill
		cell.font = header_font
		cell.alignment = Alignment(horizontal="center", vertical="center")

	for row in rows:
		built = build_cruce_row(row, datos_by_candidato.get(row.get("name")))
		ws.append([built.get(column.key, "") for column in CRUCE_COLUMNS])

	for idx in range(1, len(headers) + 1):
		ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = 16

	output = BytesIO()
	wb.save(output)
	count = len(rows)

	frappe.logger("hubgh.seleccion_cruce").info({
		"user": frappe.session.user,
		"timestamp": frappe.utils.now(),
		"count": count,
		"filters": filters,
	})

	return {
		"filename": f"cruce_seleccion_{frappe.utils.nowdate()}.xlsx",
		"content_b64": base64.b64encode(output.getvalue()).decode("ascii"),
		"count": count,
	}
