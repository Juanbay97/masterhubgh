# Copyright (c) 2026, Antigravity and contributors
# For license information, please see license.txt

"""TDD — Slice 2: role gate, bulk readers, dedupe, list/export endpoints, audit."""

import base64
import logging
from io import BytesIO
from unittest.mock import patch

import frappe
import openpyxl
from frappe.tests.utils import FrappeTestCase
from hubgh.hubgh.disponibilidad_service import (
	AUDIT_LOGGER_NAME,
	DISPONIBILIDAD_COLUMNS,
	DISPONIBILIDAD_ROLES,
	_base_candidate_filters,
	_base_hired_filters,
	_build_row_set,
	_dedupe_by_cedula,
	_read_disponibilidad_rows,
	export_disponibilidad_xlsx,
	list_disponibilidad,
)


def _fake_get_all(doctype, filters=None, fields=None, **kwargs):
	filters = filters or {}
	if doctype == "Ficha Empleado":
		return [frappe._dict({
			"name": "1001", "cedula": "1001", "nombres": "Ana", "apellidos": "Ruiz",
			"pdv": "PDV-07", "estado": "Activo", "fecha_ingreso": "2026-03-10",
		})]
	if doctype == "Candidato" and fields == ["name", "numero_documento"]:
		return [frappe._dict({"name": "CAND-1001", "numero_documento": "1001"})]
	if doctype == "Candidato":
		return [
			frappe._dict({
				"name": "1002", "numero_documento": "1002", "nombres": "Luis", "apellidos": "",
				"primer_apellido": "Gomez", "segundo_apellido": "Diaz", "pdv_destino": "PDV-01",
				"fecha_tentativa_ingreso": "2026-03-12", "estado_proceso": "En Afiliación",
			}),
			frappe._dict({
				"name": "1003", "numero_documento": "1003", "nombres": "Rechazo", "apellidos": "Test",
				"primer_apellido": "", "segundo_apellido": "", "pdv_destino": "PDV-01",
				"fecha_tentativa_ingreso": "2026-03-01", "estado_proceso": "Rechazado",
			}),
		]
	if doctype == "Candidato Disponibilidad":
		return [{"parent": "CAND-1001", "dia": "Lunes", "hora_inicio": "08:00:00", "hora_fin": "17:00:00", "idx": 1}]
	return []


class TestDisponibilidadRolesGate(FrappeTestCase):
	def test_roles_tuple_is_exact_and_all_allowed(self):
		self.assertEqual(
			DISPONIBILIDAD_ROLES,
			("GH - RRLL", "Gerente GH", "Gestión Humana", "HR Labor Relations", "Relaciones Laborales Jefe", "System Manager"),
		)

	def test_denies_out_of_set_role_before_any_domain_read_both_endpoints(self):
		domain = {"Ficha Empleado", "Candidato", "Candidato Disponibilidad"}
		for endpoint in (list_disponibilidad, export_disponibilidad_xlsx):
			with self.subTest(endpoint=endpoint.__name__):
				with (
					patch("hubgh.hubgh.disponibilidad_service.user_has_any_role", return_value=False),
					patch("hubgh.hubgh.disponibilidad_service.frappe.session") as mock_session,
					patch("hubgh.hubgh.disponibilidad_service.frappe.get_all") as mock_get_all,
				):
					mock_session.user = "hr.sst@homeburgers.com"
					with self.assertRaises(frappe.exceptions.ValidationError):
						endpoint()
				for call in mock_get_all.call_args_list:
					doctype = call.args[0] if call.args else call.kwargs.get("doctype")
					self.assertNotIn(doctype, domain, f"Unexpected domain read before gate: {call}")


class TestReadersAndDedupe(FrappeTestCase):
	def test_base_hired_filters_shapes_and_excludes_retirado(self):
		self.assertEqual(
			_base_hired_filters(fecha_desde="2026-03-01", fecha_hasta="2026-03-31", search="1001", pdv="PDV-01"),
			{"pdv": "PDV-01", "fecha_ingreso": ["between", ["2026-03-01", "2026-03-31"]], "cedula": ["like", "%1001%"], "estado": ["!=", "Retirado"]},
		)
		self.assertEqual(_base_hired_filters(fecha_desde="2026-03-01")["fecha_ingreso"], [">=", "2026-03-01"])
		self.assertEqual(_base_hired_filters(fecha_hasta="2026-03-31")["fecha_ingreso"], ["<=", "2026-03-31"])

	def test_base_candidate_filters_shapes(self):
		self.assertEqual(
			_base_candidate_filters(fecha_desde="2026-03-01", fecha_hasta="2026-03-31", search="1002", pdv="PDV-01"),
			{"pdv_destino": "PDV-01", "fecha_tentativa_ingreso": ["between", ["2026-03-01", "2026-03-31"]], "numero_documento": ["like", "%1002%"]},
		)

	def test_combined_listing_dedupe_availability_and_rechazado_exclusion(self):
		with (
			patch("hubgh.hubgh.disponibilidad_service.frappe.get_all", side_effect=_fake_get_all),
			patch("hubgh.hubgh.disponibilidad_service.get_punto_name_map", return_value={"PDV-07": "Chía", "PDV-01": "Bogotá"}),
		):
			rows = _build_row_set()
		self.assertEqual(len(rows), 2)  # Rechazado (1003) excluded
		by_cedula = {p["cedula"]: (p, avail) for p, avail in rows}
		hired_person, hired_avail = by_cedula["1001"]
		self.assertEqual(hired_person["vinculacion"], "Contratado")
		self.assertEqual(hired_person["punto_de_venta"], "Chía")  # AC-2: current PDV, not stale
		self.assertEqual(len(hired_avail), 1)  # joined via Q3 candidato-name lookup
		proceso_person, proceso_avail = by_cedula["1002"]
		self.assertEqual(proceso_person["nombre"], "Luis Gomez Diaz")  # ADR-5 fallback chain
		self.assertEqual(proceso_person["vinculacion"], "En proceso")
		self.assertEqual(proceso_avail, [])

	def test_hired_wins_when_present_in_both_populations(self):
		result = _dedupe_by_cedula({"1001"}, [({"cedula": "1001"}, "CAND-1001"), ({"cedula": "1004"}, "CAND-1004")])
		self.assertEqual([p["cedula"] for p, _ in result], ["1004"])

	def test_query_count_bounds_full_and_restricted_population(self):
		with patch("hubgh.hubgh.disponibilidad_service.frappe.get_all", side_effect=_fake_get_all) as mock_get_all:
			_build_row_set()
			self.assertLessEqual(mock_get_all.call_count, 5)
		with patch("hubgh.hubgh.disponibilidad_service.frappe.get_all", side_effect=_fake_get_all) as mock_get_all:
			_build_row_set(poblacion="contratados")
			self.assertLessEqual(mock_get_all.call_count, 4)


class TestAvailabilityReaderDirect(FrappeTestCase):
	def test_exact_filter_triple_and_never_uses_get_doc(self):
		with (
			patch("hubgh.hubgh.disponibilidad_service.frappe.get_all", return_value=[]) as mock_get_all,
			patch("hubgh.hubgh.disponibilidad_service.frappe.get_doc") as mock_get_doc,
		):
			_read_disponibilidad_rows(["CAND-1001", "CAND-1002"])
		mock_get_all.assert_called_once_with(
			"Candidato Disponibilidad",
			filters={"parent": ["in", ["CAND-1001", "CAND-1002"]], "parenttype": "Candidato", "parentfield": "disponibilidad"},
			fields=["parent", "dia", "hora_inicio", "hora_fin", "idx"],
			order_by="parent asc, idx asc",
		)
		mock_get_doc.assert_not_called()

	def test_regression_non_permlevel_role_still_gets_populated_availability(self):
		"""HR Labor Relations has NO permission row on Candidato; the direct read
		never calls get_doc, so availability populates regardless of role."""
		with (
			patch("hubgh.hubgh.disponibilidad_service.user_has_any_role", return_value=True),
			patch("hubgh.hubgh.disponibilidad_service.frappe.get_all", side_effect=_fake_get_all),
			patch("hubgh.hubgh.disponibilidad_service.get_punto_name_map", return_value={}),
			patch("hubgh.hubgh.disponibilidad_service.frappe.get_doc") as mock_get_doc,
		):
			rows = list_disponibilidad()
		mock_get_doc.assert_not_called()
		hired_row = next(r for r in rows if r["cedula"] == "1001")
		self.assertEqual(hired_row["lunes"], "08:00 - 17:00")


class TestExportDisponibilidadXlsx(FrappeTestCase):
	def test_zero_matches_headers_only_workbook_and_audit_logged(self):
		with (
			patch("hubgh.hubgh.disponibilidad_service.user_has_any_role", return_value=True),
			patch("hubgh.hubgh.disponibilidad_service.frappe.get_all", return_value=[]),
			patch("hubgh.hubgh.disponibilidad_service.frappe.logger") as mock_logger,
		):
			result = export_disponibilidad_xlsx()
		self.assertEqual(result["count"], 0)
		wb = openpyxl.load_workbook(BytesIO(base64.b64decode(result["content_b64"])))
		ws = wb.active
		self.assertEqual(ws.max_row, 1)
		self.assertEqual([c.value for c in ws[1]], [col.label for col in DISPONIBILIDAD_COLUMNS])
		mock_logger.assert_called_once_with(AUDIT_LOGGER_NAME, allow_site=True)
		info_call = mock_logger.return_value.info.call_args
		self.assertIsNotNone(info_call, "must audit-log even with count == 0")
		logged = info_call.args[0]
		self.assertEqual(logged["count"], 0)
		self.assertIn("user", logged)
		self.assertIn("timestamp", logged)

	def test_audit_logger_forced_to_info_and_real_handler_receives_record(self):
		captured = []

		class _CaptureHandler(logging.Handler):
			def emit(self, record):
				captured.append(record)

		logger = frappe.logger(AUDIT_LOGGER_NAME, allow_site=True)
		handler = _CaptureHandler()
		logger.addHandler(handler)
		try:
			with (
				patch("hubgh.hubgh.disponibilidad_service.user_has_any_role", return_value=True),
				patch("hubgh.hubgh.disponibilidad_service.frappe.get_all", return_value=[]),
			):
				export_disponibilidad_xlsx()
		finally:
			logger.removeHandler(handler)
		self.assertEqual(logger.level, logging.INFO)
		info_records = [r for r in captured if r.levelno == logging.INFO]
		self.assertEqual(len(info_records), 1, "real logger handler must receive the audit record")
		self.assertEqual(info_records[0].msg["count"], 0)
