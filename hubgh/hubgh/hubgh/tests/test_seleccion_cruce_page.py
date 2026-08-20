# Copyright (c) 2026, Antigravity and contributors
# For license information, please see license.txt

"""
TDD — Batch PR3, Phases 3.1-3.3: cruce tray page + Excel export endpoints.

Tests cover:
  3.2.1  RED (security) → list_cruce_candidates denies before any DB read
  3.2.3  RED  → tray excludes Rechazado always, regardless of filters
  3.2.5  RED  → empty result set returns empty payload, no error
  3.3.1  RED (security) → export_cruce_xlsx denies before any DB read
  3.3.3  RED  → export with zero matching candidates: headers-only workbook + audit log
  3.3.4  GREEN → audit log call assertion (user, timestamp, count)
  3.3.5  RED  → golden export: fully populated candidate header row matches CRUCE_COLUMNS

All tests are RED until the corresponding GREEN task lands (see tasks.md Phase 3.2/3.3).
"""

import base64
import logging
from io import BytesIO
from unittest.mock import patch

import frappe
import openpyxl
from frappe.tests.utils import FrappeTestCase

from hubgh.hubgh.seleccion_cruce_service import (
	AUDIT_LOGGER_NAME,
	CRUCE_COLUMNS,
	export_cruce_xlsx,
	list_cruce_candidates,
)


def _candidato_row(**overrides):
	base = {
		"name": "CAND-100",
		"tipo_documento": "Cedula",
		"numero_documento": "999888",
		"nombres": "Maria",
		"apellidos": "",
		"primer_apellido": "Perez",
		"segundo_apellido": "Lopez",
		"pdv_destino": "PDV-01",
		"cargo_postulado": "Asesor",
		"creation": "2026-01-01 10:00:00",
		"estado_proceso": "En documentación",
		"fecha_tentativa_ingreso": "2026-02-01",
		"fecha_nacimiento": "1995-05-20",
		"direccion": "Calle 1 # 2-3",
		"barrio": "Centro",
		"ciudad": "BOG",
		"celular": "3000000000",
		"email": "maria@test.com",
		"telefono_fijo": "6011234567",
		"grupo_sanguineo": "O+",
		"contacto_emergencia_nombre": "Luis",
		"contacto_emergencia_telefono": "3001111111",
		"tiene_alergias": 0,
		"descripcion_alergias": "",
		"talla_camisa": "M",
		"talla_pantalon": "32",
		"numero_zapatos": "38",
		"talla_delantal": "L",
		"estado_civil": "Soltero",
		"nivel_educativo_siesa": None,
		"eps_siesa": "EPS1",
		"afp_siesa": "AFP1",
		"cesantias_siesa": "CES1",
		"banco_siesa": "BANCO1",
		"numero_cuenta_bancaria": "0001",
	}
	base.update(overrides)
	return frappe._dict(base)


def _assert_no_candidate_domain_read(mock_get_all):
	"""Assert no call queried Candidato/Datos Contratacion — ignores frappe-internal
	calls (e.g. the Translation lookup triggered by frappe.throw's own i18n)."""
	domain_doctypes = {"Candidato", "Datos Contratacion"}
	for call in mock_get_all.call_args_list:
		doctype = call.args[0] if call.args else call.kwargs.get("doctype")
		assert doctype not in domain_doctypes, f"Unexpected candidate-domain DB read before gate: {call}"


# ---------------------------------------------------------------------------
# 3.2.1 — security gate before any DB read: list_cruce_candidates
# ---------------------------------------------------------------------------


class TestListCruceCandidatesGate(FrappeTestCase):
	def test_denies_non_hr_read_role_before_db_read(self):
		with (
			patch("hubgh.hubgh.seleccion_cruce_service.user_has_any_role", return_value=False),
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.session") as mock_session,
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.get_all") as mock_get_all,
		):
			mock_session.user = "hr.sst@homeburgers.com"
			with self.assertRaises(frappe.exceptions.ValidationError):
				list_cruce_candidates()
		_assert_no_candidate_domain_read(mock_get_all)


class TestListCruceCandidatesFilters(FrappeTestCase):
	"""3.2.3/3.2.5 — Rechazado exclusion + empty payload."""

	def test_excludes_rechazado_always_even_when_filter_requests_it(self):
		rows = [
			_candidato_row(name="CAND-100", estado_proceso="En documentación"),
			_candidato_row(name="CAND-101", estado_proceso="Rechazado"),
		]
		with (
			patch("hubgh.hubgh.seleccion_cruce_service.user_has_any_role", return_value=True),
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.get_all", return_value=rows) as mock_get_all,
			patch("hubgh.hubgh.seleccion_cruce_service.get_punto_name_map", return_value={}),
		):
			result = list_cruce_candidates(estado="Rechazado")
		self.assertEqual(result, [])
		mock_get_all.assert_called_once()

	def test_excludes_rechazado_from_unfiltered_listing(self):
		rows = [
			_candidato_row(name="CAND-100", estado_proceso="En documentación"),
			_candidato_row(name="CAND-101", estado_proceso="Rechazado"),
		]
		with (
			patch("hubgh.hubgh.seleccion_cruce_service.user_has_any_role", return_value=True),
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.get_all", return_value=rows),
			patch("hubgh.hubgh.seleccion_cruce_service.get_punto_name_map", return_value={}),
		):
			result = list_cruce_candidates()
		self.assertEqual([r["name"] for r in result], ["CAND-100"])

	def test_empty_result_set_returns_empty_payload_no_error(self):
		with (
			patch("hubgh.hubgh.seleccion_cruce_service.user_has_any_role", return_value=True),
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.get_all", return_value=[]),
			patch("hubgh.hubgh.seleccion_cruce_service.get_punto_name_map", return_value={}),
		):
			result = list_cruce_candidates(pdv="PDV-NOEXISTE")
		self.assertEqual(result, [])


# ---------------------------------------------------------------------------
# 3.3.1 — security gate before any DB read: export_cruce_xlsx
# ---------------------------------------------------------------------------


class TestExportCruceXlsxGate(FrappeTestCase):
	def test_denies_non_hr_read_role_before_db_read(self):
		with (
			patch("hubgh.hubgh.seleccion_cruce_service.user_has_any_role", return_value=False),
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.session") as mock_session,
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.get_all") as mock_get_all,
		):
			mock_session.user = "hr.sst@homeburgers.com"
			with self.assertRaises(frappe.exceptions.ValidationError):
				export_cruce_xlsx()
		_assert_no_candidate_domain_read(mock_get_all)


# ---------------------------------------------------------------------------
# 3.3.3/3.3.4 — zero-match export: headers-only workbook, still audit-logged
# ---------------------------------------------------------------------------


class TestExportCruceXlsxZeroMatches(FrappeTestCase):
	def test_zero_candidates_returns_headers_only_workbook_and_logs(self):
		with (
			patch("hubgh.hubgh.seleccion_cruce_service.user_has_any_role", return_value=True),
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.get_all", return_value=[]),
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.logger") as mock_logger,
		):
			result = export_cruce_xlsx()

		self.assertEqual(result["count"], 0)
		wb = openpyxl.load_workbook(BytesIO(base64.b64decode(result["content_b64"])))
		ws = wb.active
		self.assertEqual(ws.max_row, 1)  # only the header row
		self.assertEqual([c.value for c in ws[1]], [col.label for col in CRUCE_COLUMNS])

		mock_logger.assert_called_once_with("hubgh.seleccion_cruce", allow_site=True)
		info_call = mock_logger.return_value.info.call_args
		self.assertIsNotNone(info_call, "export_cruce_xlsx must audit-log even with count == 0")
		logged = info_call.args[0]
		self.assertEqual(logged["count"], 0)
		self.assertIn("user", logged)
		self.assertIn("timestamp", logged)


class TestExportCruceXlsxAuditLogRealHandler(FrappeTestCase):
	"""Remediation (verify finding): frappe.logger() defaults to effective
	level ERROR outside dev server, so a bare .info() call is silently dropped
	in production. Uses the REAL logger (not mocked) with a capture handler."""

	def test_audit_logger_forced_to_info_and_real_handler_receives_record(self):
		captured = []

		class _CaptureHandler(logging.Handler):
			def emit(self, record):
				captured.append(record)

		logger = frappe.logger(AUDIT_LOGGER_NAME, allow_site=True)
		capture_handler = _CaptureHandler()
		logger.addHandler(capture_handler)
		try:
			with (
				patch("hubgh.hubgh.seleccion_cruce_service.user_has_any_role", return_value=True),
				patch("hubgh.hubgh.seleccion_cruce_service.frappe.get_all", return_value=[]),
			):
				export_cruce_xlsx()
		finally:
			logger.removeHandler(capture_handler)

		self.assertEqual(logger.level, logging.INFO)
		info_records = [r for r in captured if r.levelno == logging.INFO]
		self.assertEqual(len(info_records), 1, "real logger handler must receive the audit record")
		logged = info_records[0].msg
		self.assertEqual(logged["count"], 0)
		self.assertIn("user", logged)
		self.assertIn("timestamp", logged)
# ---------------------------------------------------------------------------
# 3.3.5 — golden export: header row matches the 49-column template exactly
# ---------------------------------------------------------------------------


class TestExportCruceXlsxGolden(FrappeTestCase):
	def test_populated_candidate_export_header_matches_template(self):
		row = _candidato_row()
		with (
			patch("hubgh.hubgh.seleccion_cruce_service.user_has_any_role", return_value=True),
			patch("hubgh.hubgh.seleccion_cruce_service.frappe.get_all", side_effect=[[row], []]),
		):
			result = export_cruce_xlsx()

		wb = openpyxl.load_workbook(BytesIO(base64.b64decode(result["content_b64"])))
		ws = wb.active
		header_row = [c.value for c in ws[1]]
		self.assertEqual(header_row, [col.label for col in CRUCE_COLUMNS])
		self.assertEqual(result["count"], 1)
		# One data row present beyond the header.
		self.assertEqual(ws.max_row, 2)
